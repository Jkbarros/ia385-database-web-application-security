#!/usr/bin/env python3
"""Convert archive/canvas-submissions/*.docx to docs/*.md."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
ARCHIVE = ROOT / "archive" / "canvas-submissions"
DOCS = ROOT / "docs"

SLUGS: dict[str, str] = {
    "Lab 1 Final.docx": "lab-1-lab-build.md",
    "Lab 2 Final.docx": "lab-2-mysql-install.md",
    "Assignment 3.docx": "assignment-3-sql-server.md",
    "Assignment 4 Install Oracle 1.docx": "assignment-4-oracle-install.md",
    "Assignment 5.docx": "assignment-5-users-roles.md",
    "Assignment 6.docx": "assignment-6-sql-queries.md",
}


def table_to_md(table) -> list[str]:
    rows: list[list[str]] = []
    for row in table.rows:
        rows.append([cell.text.strip().replace("\n", " ") for cell in row.cells])
    if not rows:
        return []
    lines = [
        "| " + " | ".join(rows[0]) + " |",
        "| " + " | ".join("---" for _ in rows[0]) + " |",
    ]
    for row in rows[1:]:
        lines.append("| " + " | ".join(row) + " |")
    return lines


def convert_docx(src: Path, dest: Path, title: str) -> None:
    doc = Document(src)
    rel = src.relative_to(ROOT).as_posix().replace(" ", "%20")
    lines: list[str] = [
        f"# {title}",
        "",
        f"**Source:** [`{src.name}`](../{rel})  ",
        "**Course:** IA 385 - Database and Web Application Security (EMU)  ",
        "**Author:** Julio Barros",
        "",
    ]
    for para in doc.paragraphs:
        t = para.text.strip()
        if not t:
            lines.append("")
            continue
        style = (para.style.name or "").lower()
        if "heading 1" in style:
            lines.extend(["", f"## {t}", ""])
        elif "heading 2" in style:
            lines.extend(["", f"### {t}", ""])
        elif "heading 3" in style:
            lines.extend(["", f"#### {t}", ""])
        elif "list" in style:
            lines.append(f"- {t}")
        else:
            if t.startswith("SET ") or t.startswith("ALTER ") or t.startswith("SELECT "):
                lines.extend(["", "```sql", t, "```", ""])
            else:
                lines.append(t)
                lines.append("")

    for table in doc.tables:
        lines.extend(["", *table_to_md(table), ""])

    body = "\n".join(lines)
    body = re.sub(r"\n{3,}", "\n\n", body).strip() + "\n"
    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_text(body, encoding="utf-8")


def main() -> None:
    DOCS.mkdir(parents=True, exist_ok=True)
    for name, slug in SLUGS.items():
        src = ARCHIVE / name
        if not src.exists():
            print(f"skip missing {name}")
            continue
        title = name.removesuffix(".docx")
        convert_docx(src, DOCS / slug, title)
        print(f"wrote docs/{slug}")


if __name__ == "__main__":
    main()
